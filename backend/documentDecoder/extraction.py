"""
Extracts plain text from an uploaded PDF, DOCX, or TXT file.

Kept isolated from main.py so the extraction strategy (and future OCR
support) can change without touching request handling.
"""
from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument


class ExtractionError(Exception):
    """Raised when text cannot be extracted from the file."""


class ScannedPdfError(ExtractionError):
    """Raised when a PDF appears to contain no extractable text (likely scanned/image-only)."""


MAX_CHARS = 400_000  # guard against extremely large documents blowing up the AI call


def _extract_pdf(path: Path) -> str:
    chunks = []
    try:
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) == 0:
                raise ExtractionError("The PDF has no pages.")
            for page in pdf.pages:
                text = page.extract_text() or ""
                chunks.append(text)

                # Also pull table content, since extract_text() can miss tables.
                for table in page.extract_tables() or []:
                    for row in table:
                        chunks.append(" | ".join(cell or "" for cell in row))
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not read this PDF: {exc}") from exc

    combined = "\n".join(c for c in chunks if c).strip()
    if not combined:
        raise ScannedPdfError(
            "No extractable text was found in this PDF. It may be a scanned or "
            "image-only document. Please upload a text-based PDF, or a DOCX/TXT version."
        )
    return combined


def _extract_docx(path: Path) -> str:
    try:
        doc = DocxDocument(path)
    except Exception as exc:
        raise ExtractionError(f"Could not read this DOCX file: {exc}") from exc

    parts = [p.text for p in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    combined = "\n".join(p for p in parts if p and p.strip()).strip()
    if not combined:
        raise ExtractionError("This DOCX file appears to be empty.")
    return combined


def _extract_txt(path: Path) -> str:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        raise ExtractionError(f"Could not read this text file: {exc}") from exc

    if not text.strip():
        raise ExtractionError("This text file appears to be empty.")
    return text


def extract_text(path: Path, filename: str) -> str:
    """Dispatch extraction by file extension. Returns extracted text or raises ExtractionError."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    elif suffix == ".txt":
        text = _extract_txt(path)
    else:
        raise ExtractionError(
            f"Unsupported file type '{suffix}'. Please upload a PDF, DOCX, or TXT file."
        )

    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n[Document truncated for length.]"

    return text
